import json
import shutil
from pathlib import Path

import yaml
from openpyxl import load_workbook

from risk_model_workbench.cli import main
from risk_model_workbench.planning import create_execution_plan
from risk_model_workbench.project.create import create_project
from risk_model_workbench.reporting.excel_report import generate_excel_report
from risk_model_workbench.request import parse_model_request


def test_generic_project_templates_do_not_default_to_gcard(tmp_path):
    project = create_project(
        tmp_path,
        name="generic-model",
        display_name="Generic Model",
        scenario="Generic Scenario",
    )

    for path in [
        project / "configs" / "evaluate.yaml",
        project / "configs" / "train.yaml",
        project / "queries" / "02_split_summary.sql",
        project / "queries" / "03_segment_zc_summary.sql",
    ]:
        assert "gcard" not in path.read_text(encoding="utf-8").lower()

    evaluate = yaml.safe_load((project / "configs" / "evaluate.yaml").read_text(encoding="utf-8"))
    assert evaluate["evaluation"]["score_columns"] == ["model_score"]
    assert evaluate["evaluation"]["score_labels"] == {"model_score": "本轮模型"}

    train = yaml.safe_load((project / "configs" / "train.yaml").read_text(encoding="utf-8"))
    assert train["input"]["historical_score_columns"] == []


def test_generic_execution_plan_and_compare_do_not_default_to_gcard(tmp_path):
    project = create_project(
        tmp_path,
        name="generic-model",
        display_name="Generic Model",
        scenario="Generic Scenario",
    )
    request_doc = parse_model_request(project / "requests" / "model_request_template.md")

    plan = create_execution_plan(request_doc, project)
    compare_task = next(task for task in plan["tasks"] if task["task_id"] == "compare_final")

    assert "--champion" not in compare_task["command"]["args"]
    assert "gcard_v6" not in compare_task["command"]["args"]

    assert main(["run", "init", "--project", str(project), "--workflow", "full_modeling", "--run-id", "run1"]) == 0
    assert main(["compare", "--project", str(project), "--run-id", "run1"]) == 0
    payload = json.loads((project / "runs" / "run1" / "evaluation" / "champion_challenger.json").read_text(encoding="utf-8"))
    assert payload["status"] == "skipped"
    assert payload["reason"] == "no champion configured"


def test_generic_report_does_not_emit_gcard_defaults(tmp_path):
    project = create_project(
        tmp_path,
        name="generic-model",
        display_name="Generic Model",
        scenario="Generic Scenario",
    )
    run_dir = project / "runs" / "run1"
    eval_dir = run_dir / "evaluation"
    train_dir = run_dir / "modeling" / "main_lgbm"
    input_dir = run_dir / "modeling_input"
    feature_dir = run_dir / "feature_selection"
    for directory in [eval_dir, train_dir, input_dir, feature_dir, run_dir / "reports"]:
        directory.mkdir(parents=True, exist_ok=True)
    (eval_dir / "evaluation_summary.json").write_text(
        json.dumps({"score_columns_evaluated": ["model_score"], "splits_evaluated": ["DEV"], "n_total_samples": 10}) + "\n",
        encoding="utf-8",
    )
    (train_dir / "metrics_train_valid.json").write_text(
        json.dumps({"valid_auc": 0.7, "valid_ks": 0.3, "auc_gap": 0.01}) + "\n",
        encoding="utf-8",
    )
    (train_dir / "run_config.json").write_text(
        json.dumps({"label_column": "target", "split_column": "final_flag", "algorithm": "lightgbm"}) + "\n",
        encoding="utf-8",
    )
    (feature_dir / "feature_stage_summary.json").write_text(json.dumps({"final_training_features": 3}) + "\n", encoding="utf-8")

    output_path = run_dir / "reports" / "model_report.xlsx"
    generate_excel_report(
        eval_dir=eval_dir,
        train_dir=train_dir,
        input_dir=input_dir,
        feature_dir=feature_dir,
        output_path=output_path,
        project_dir=project,
    )

    report_text = output_path.with_name("model_report.md").read_text(encoding="utf-8")
    report_html = output_path.with_name("model_report.html").read_text(encoding="utf-8")
    missing_text = output_path.with_name("model_report_missing_results.md").read_text(encoding="utf-8")
    workbook = load_workbook(output_path)
    assert "gcard" not in report_text.lower()
    assert "gcard" not in report_html.lower()
    assert "G卡" not in report_html
    assert "gcard" not in missing_text.lower()
    assert "Summary" not in workbook.sheetnames
    assert "# Generic Model模型报告" in report_text


def test_gcard_execution_plan_still_uses_configured_champion():
    request_doc = parse_model_request("projects/2026-05-fujie-gcard-v1/requests/model_request_template.md")

    plan = create_execution_plan(request_doc, "projects/2026-05-fujie-gcard-v1")
    compare_task = next(task for task in plan["tasks"] if task["task_id"] == "compare_final")

    args = compare_task["command"]["args"]
    assert args.count("--champion") == 4
    assert args[-8:] == [
        "--champion",
        "gcard_v2",
        "--champion",
        "gcard_v4",
        "--champion",
        "gcard_v5",
        "--champion",
        "gcard_v6",
    ]


def test_gcard_compare_with_explicit_champion_still_supported():
    project = Path("projects/2026-05-fujie-gcard-v1")
    run_id = "pytest_gcard_compare_compat"
    run_dir = project / "runs" / run_id
    try:
        assert main(["run", "init", "--project", str(project), "--workflow", "full_modeling", "--run-id", run_id, "--force"]) == 0
        assert main(["compare", "--project", str(project), "--run-id", run_id, "--champion", "gcard_v6"]) == 0
        payload = json.loads((run_dir / "evaluation" / "champion_challenger.json").read_text(encoding="utf-8"))
        assert payload["status"] == "scaffold"
        assert payload["champion"] == "gcard_v6"
    finally:
        shutil.rmtree(run_dir, ignore_errors=True)
